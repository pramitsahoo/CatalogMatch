"""
Export Product Matching Results to Excel with Images
=====================================================
Creates a comprehensive Excel report with embedded images for HR review.
"""

import pandas as pd
import numpy as np
from pathlib import Path
from PIL import Image
from io import BytesIO
import base64
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
import tempfile
import os

# Configuration
IMAGES_DIR = Path("dataset/images")
PRODUCTS_CSV = Path("dataset/products.csv")
QUERIES_CSV = Path("dataset/queries.csv")
OUTPUT_EXCEL = Path("product_matching_results.xlsx")
THUMBNAIL_SIZE = (80, 80)

def resize_image_for_excel(image_path: Path, size: tuple = THUMBNAIL_SIZE) -> str:
    """Resize image and return temp file path."""
    try:
        img = Image.open(image_path)
        img.thumbnail(size, Image.Resampling.LANCZOS)

        # Convert to RGB if necessary
        if img.mode in ('RGBA', 'P'):
            img = img.convert('RGB')

        # Save to temp file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.jpg')
        img.save(temp_file.name, 'JPEG', quality=85)
        return temp_file.name
    except Exception as e:
        print(f"Error processing image {image_path}: {e}")
        return None


def create_excel_report(
    results_df: pd.DataFrame,
    products_df: pd.DataFrame,
    queries_df: pd.DataFrame,
    metrics: dict,
    output_path: Path = OUTPUT_EXCEL,
    include_images: bool = True
):
    """
    Create comprehensive Excel report with images.

    Args:
        results_df: DataFrame with columns [query_id, has_match, ground_truth_db_id,
                    predicted_match, best_score, top1_db_id, is_correct]
        products_df: Products database
        queries_df: Query dataset
        metrics: Dictionary with all evaluation metrics
        output_path: Output Excel file path
        include_images: Whether to embed images
    """
    wb = Workbook()

    # ========== Sheet 1: Summary Metrics ==========
    ws_summary = wb.active
    ws_summary.title = "Summary"

    # Title
    ws_summary['A1'] = "Two-Stage Multimodal Product Matching - Results Summary"
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary.merge_cells('A1:F1')

    # Pipeline Description
    ws_summary['A3'] = "Pipeline Architecture"
    ws_summary['A3'].font = Font(bold=True, size=12)
    ws_summary['A4'] = "Stage 1 (Retrieval): Qwen3-VL-Embedding-2B + ChromaDB"
    ws_summary['A5'] = "Stage 2 (Reranking): Qwen3-VL-Reranker-2B"

    # Dataset Stats
    ws_summary['A7'] = "Dataset Statistics"
    ws_summary['A7'].font = Font(bold=True, size=12)
    ws_summary['A8'] = f"Products in DB: {len(products_df)}"
    ws_summary['A9'] = f"Total Queries: {len(queries_df)}"
    ws_summary['A10'] = f"Queries with Match: {len(queries_df[queries_df['has_match'] == True])}"
    ws_summary['A11'] = f"Queries without Match: {len(queries_df[queries_df['has_match'] == False])}"

    # Retrieval Metrics
    ws_summary['A13'] = "Stage 1: Retrieval Metrics"
    ws_summary['A13'].font = Font(bold=True, size=12)
    row = 14
    for key, value in metrics.get('retrieval', {}).items():
        ws_summary[f'A{row}'] = key
        ws_summary[f'B{row}'] = f"{value:.4f}" if isinstance(value, float) else str(value)
        row += 1

    # Reranking Metrics
    ws_summary[f'A{row+1}'] = "Stage 2: Reranking Metrics"
    ws_summary[f'A{row+1}'].font = Font(bold=True, size=12)
    row += 2
    for key, value in metrics.get('reranking', {}).items():
        ws_summary[f'A{row}'] = key
        ws_summary[f'B{row}'] = f"{value:.4f}" if isinstance(value, float) else str(value)
        row += 1

    # Binary Classification Metrics
    ws_summary[f'A{row+1}'] = "Binary Classification"
    ws_summary[f'A{row+1}'].font = Font(bold=True, size=12)
    row += 2
    for key, value in metrics.get('classification', {}).items():
        ws_summary[f'A{row}'] = key
        ws_summary[f'B{row}'] = f"{value:.4f}" if isinstance(value, float) else str(value)
        row += 1

    # Production Thresholds (Human-in-the-Loop)
    if 'production_thresholds' in metrics:
        ws_summary[f'A{row+1}'] = "Production Thresholds (Human-in-the-Loop)"
        ws_summary[f'A{row+1}'].font = Font(bold=True, size=12)
        row += 2
        for key, value in metrics.get('production_thresholds', {}).items():
            ws_summary[f'A{row}'] = key
            ws_summary[f'B{row}'] = f"{value:.4f}" if isinstance(value, float) else str(value)
            row += 1

        # Decision zone explanation
        ws_summary[f'A{row+1}'] = "Decision Zones:"
        ws_summary[f'A{row+2}'] = "  • Auto-MERGE: score ≥ upper_threshold (confident match)"
        ws_summary[f'A{row+3}'] = "  • Human-REVIEW: between thresholds (uncertain)"
        ws_summary[f'A{row+4}'] = "  • Auto-REJECT: score < lower_threshold (confident no-match)"
        row += 5

    # ========== Sheet 2: Detailed Results ==========
    ws_results = wb.create_sheet("Detailed Results")

    # Headers
    headers = [
        "Query ID", "Query Name", "Query Image",
        "Has Match", "Ground Truth ID", "Ground Truth Name", "Ground Truth Image",
        "Predicted Match", "Top-1 Match ID", "Top-1 Match Name", "Top-1 Image",
        "Rerank Score", "Decision Zone", "Result", "Status"
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws_results.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Column widths
    col_widths = [15, 40, 15, 12, 15, 40, 15, 15, 15, 40, 15, 12, 15, 10, 12]
    for i, width in enumerate(col_widths, 1):
        ws_results.column_dimensions[get_column_letter(i)].width = width

    # Colors for results
    correct_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    incorrect_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Data rows
    temp_files = []  # Track temp files for cleanup
    current_row = 2
    row_height = 60 if include_images else 15

    for idx, row in results_df.iterrows():
        query_row = queries_df[queries_df['query_id'] == row['query_id']].iloc[0]

        # Query info
        ws_results.cell(row=current_row, column=1, value=row['query_id'])
        ws_results.cell(row=current_row, column=2, value=query_row['name'][:50] + "..." if len(str(query_row['name'])) > 50 else query_row['name'])

        # Query image
        if include_images:
            query_img_path = IMAGES_DIR / query_row['image']
            if query_img_path.exists():
                temp_path = resize_image_for_excel(query_img_path)
                if temp_path:
                    temp_files.append(temp_path)
                    img = XLImage(temp_path)
                    ws_results.add_image(img, f"C{current_row}")

        ws_results.cell(row=current_row, column=4, value="Yes" if row['has_match'] else "No")

        # Ground truth info
        if row['has_match']:
            gt_id = row['ground_truth_db_id']
            gt_row = products_df[products_df['db_id'] == gt_id].iloc[0]
            ws_results.cell(row=current_row, column=5, value=gt_id)
            ws_results.cell(row=current_row, column=6, value=gt_row['name'][:50] + "..." if len(str(gt_row['name'])) > 50 else gt_row['name'])

            if include_images:
                gt_img_path = IMAGES_DIR / gt_row['image']
                if gt_img_path.exists():
                    temp_path = resize_image_for_excel(gt_img_path)
                    if temp_path:
                        temp_files.append(temp_path)
                        img = XLImage(temp_path)
                        ws_results.add_image(img, f"G{current_row}")
        else:
            ws_results.cell(row=current_row, column=5, value="-")
            ws_results.cell(row=current_row, column=6, value="No match in DB")

        # Prediction info
        ws_results.cell(row=current_row, column=8, value="Yes" if row['predicted_match'] else "No")

        if row.get('top1_db_id'):
            top1_id = row['top1_db_id']
            top1_row = products_df[products_df['db_id'] == top1_id]
            if not top1_row.empty:
                top1_row = top1_row.iloc[0]
                ws_results.cell(row=current_row, column=9, value=top1_id)
                ws_results.cell(row=current_row, column=10, value=top1_row['name'][:50] + "..." if len(str(top1_row['name'])) > 50 else top1_row['name'])

                if include_images:
                    top1_img_path = IMAGES_DIR / top1_row['image']
                    if top1_img_path.exists():
                        temp_path = resize_image_for_excel(top1_img_path)
                        if temp_path:
                            temp_files.append(temp_path)
                            img = XLImage(temp_path)
                            ws_results.add_image(img, f"K{current_row}")

        ws_results.cell(row=current_row, column=12, value=f"{row['best_score']:.3f}")

        # Decision Zone (Human-in-the-Loop)
        decision_zone = row.get('decision_zone', 'N/A')
        zone_cell = ws_results.cell(row=current_row, column=13, value=decision_zone if decision_zone else "N/A")

        # Zone-specific colors
        zone_colors = {
            "Auto-MERGE": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),  # Light blue
            "Human-REVIEW": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),  # Light yellow
            "Auto-REJECT": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),  # Light orange
        }
        if decision_zone in zone_colors:
            zone_cell.fill = zone_colors[decision_zone]

        # Result status
        status = row.get('status', 'Unknown')
        result_cell = ws_results.cell(row=current_row, column=14, value=status)

        is_correct = row.get('is_correct', False)
        status_cell = ws_results.cell(row=current_row, column=15, value="Correct" if is_correct else "Incorrect")

        # Color coding for correctness (apply to all columns except Decision Zone)
        fill = correct_fill if is_correct else incorrect_fill
        for col in range(1, 16):
            if col != 13:  # Skip Decision Zone column (keep its own color)
                ws_results.cell(row=current_row, column=col).fill = fill

        ws_results.row_dimensions[current_row].height = row_height
        current_row += 1

    # ========== Sheet 3: Error Analysis ==========
    ws_errors = wb.create_sheet("Error Analysis")

    # False Positives
    ws_errors['A1'] = "False Positives (Predicted match but no actual match)"
    ws_errors['A1'].font = Font(bold=True, size=12)

    fp_df = results_df[(results_df['predicted_match'] == True) & (results_df['has_match'] == False)]
    if not fp_df.empty:
        row = 3
        ws_errors['A2'] = "Query ID"
        ws_errors['B2'] = "Query Name"
        ws_errors['C2'] = "Best Score"
        ws_errors['D2'] = "Top-1 Match"
        for col in ['A', 'B', 'C', 'D']:
            ws_errors[f'{col}2'].font = Font(bold=True)

        for _, fp_row in fp_df.iterrows():
            query_row = queries_df[queries_df['query_id'] == fp_row['query_id']].iloc[0]
            ws_errors[f'A{row}'] = fp_row['query_id']
            ws_errors[f'B{row}'] = query_row['name'][:60]
            ws_errors[f'C{row}'] = f"{fp_row['best_score']:.3f}"
            ws_errors[f'D{row}'] = fp_row.get('top1_db_id', 'N/A')
            row += 1

    # False Negatives
    fn_start = row + 2
    ws_errors[f'A{fn_start}'] = "False Negatives (Has match but not predicted)"
    ws_errors[f'A{fn_start}'].font = Font(bold=True, size=12)

    fn_df = results_df[(results_df['predicted_match'] == False) & (results_df['has_match'] == True)]
    if not fn_df.empty:
        row = fn_start + 2
        ws_errors[f'A{fn_start+1}'] = "Query ID"
        ws_errors[f'B{fn_start+1}'] = "Query Name"
        ws_errors[f'C{fn_start+1}'] = "Best Score"
        ws_errors[f'D{fn_start+1}'] = "Ground Truth"
        for col in ['A', 'B', 'C', 'D']:
            ws_errors[f'{col}{fn_start+1}'].font = Font(bold=True)

        for _, fn_row in fn_df.iterrows():
            query_row = queries_df[queries_df['query_id'] == fn_row['query_id']].iloc[0]
            ws_errors[f'A{row}'] = fn_row['query_id']
            ws_errors[f'B{row}'] = query_row['name'][:60]
            ws_errors[f'C{row}'] = f"{fn_row['best_score']:.3f}"
            ws_errors[f'D{row}'] = fn_row.get('ground_truth_db_id', 'N/A')
            row += 1

    # ========== Sheet 4: Retrieval vs Reranking Comparison ==========
    ws_comparison = wb.create_sheet("Retrieval vs Reranking")

    ws_comparison['A1'] = "Stage Comparison: Retrieval (Fast) vs Reranking (Accurate)"
    ws_comparison['A1'].font = Font(bold=True, size=14)
    ws_comparison.merge_cells('A1:E1')

    # Comparison table
    comparison_headers = ["Metric", "Retrieval", "Reranking", "Improvement", "Notes"]
    for col, header in enumerate(comparison_headers, 1):
        cell = ws_comparison.cell(row=3, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    # Data
    comparison_data = [
        ("Recall@1", metrics['retrieval'].get('Recall@1', 0), metrics['reranking'].get('Rerank_Recall@1', 0)),
        ("Recall@3", metrics['retrieval'].get('Recall@3', 0), metrics['reranking'].get('Rerank_Recall@3', 0)),
        ("Recall@5", metrics['retrieval'].get('Recall@5', 0), metrics['reranking'].get('Rerank_Recall@5', 0)),
        ("MRR", metrics['retrieval'].get('MRR', 0), metrics['reranking'].get('Rerank_MRR', 0)),
    ]

    row = 4
    for metric, ret_val, rerank_val in comparison_data:
        ws_comparison.cell(row=row, column=1, value=metric)
        ws_comparison.cell(row=row, column=2, value=f"{ret_val:.4f}")
        ws_comparison.cell(row=row, column=3, value=f"{rerank_val:.4f}")
        improvement = rerank_val - ret_val
        ws_comparison.cell(row=row, column=4, value=f"+{improvement:.4f}" if improvement >= 0 else f"{improvement:.4f}")

        # Notes
        if metric == "Recall@1":
            ws_comparison.cell(row=row, column=5, value="Key metric: correct match at top position")
        elif metric == "MRR":
            ws_comparison.cell(row=row, column=5, value="Mean Reciprocal Rank for ranking quality")
        row += 1

    # Key insights
    ws_comparison[f'A{row+2}'] = "Key Insights:"
    ws_comparison[f'A{row+2}'].font = Font(bold=True, size=12)
    ws_comparison[f'A{row+3}'] = "1. Retrieval provides fast initial filtering (top-K candidates from entire DB)"
    ws_comparison[f'A{row+4}'] = "2. Reranking improves precision by cross-encoding query-document pairs"
    ws_comparison[f'A{row+5}'] = "3. Two-stage approach balances speed and accuracy for production use"

    # Save workbook
    wb.save(output_path)
    print(f"Excel report saved to: {output_path}")

    # Cleanup temp files
    for temp_file in temp_files:
        try:
            os.unlink(temp_file)
        except:
            pass

    return output_path


def prepare_results_dataframe(
    reranked_results: list,
    threshold: float,
    upper_threshold: float = None,
    lower_threshold: float = None
) -> pd.DataFrame:
    """
    Convert reranked results to a DataFrame for Excel export.

    Args:
        reranked_results: List of result dictionaries from pipeline
        threshold: Classification threshold (for binary decision)
        upper_threshold: Score above this → Auto-merge (optional, for 3-zone system)
        lower_threshold: Score below this → Auto-reject (optional, for 3-zone system)

    Returns:
        DataFrame with result columns including decision_zone
    """
    rows = []
    for r in reranked_results:
        predicted_match = r['best_score'] >= threshold

        # Determine correctness
        if r['has_match']:
            # Should find the correct match AND it must be the right product
            if predicted_match and r['reranked'] and r['reranked'][0]['db_id'] == r.get('ground_truth_db_id'):
                is_correct = True
                status = "TP"
            else:
                is_correct = False
                status = "FN"
        else:
            # Should NOT predict a match
            is_correct = not predicted_match
            status = "TN" if is_correct else "FP"

        # Determine decision zone (Human-in-the-Loop)
        if upper_threshold is not None and lower_threshold is not None:
            if r['best_score'] >= upper_threshold:
                decision_zone = "Auto-MERGE"
            elif r['best_score'] < lower_threshold:
                decision_zone = "Auto-REJECT"
            else:
                decision_zone = "Human-REVIEW"
        else:
            decision_zone = None

        rows.append({
            'query_id': r['query_id'],
            'has_match': r['has_match'],
            'ground_truth_db_id': r.get('ground_truth_db_id'),
            'predicted_match': predicted_match,
            'best_score': r['best_score'],
            'top1_db_id': r['reranked'][0]['db_id'] if r['reranked'] else None,
            'is_correct': is_correct,
            'status': status,
            'decision_zone': decision_zone
        })

    return pd.DataFrame(rows)


if __name__ == "__main__":
    # Example usage - this would be called from the notebook
    print("This module should be imported and used from the notebook.")
    print("Example:")
    print("  from export_results_to_excel import create_excel_report, prepare_results_dataframe")
    print("  results_df = prepare_results_dataframe(reranked_results, threshold)")
    print("  create_excel_report(results_df, products_df, queries_df, metrics)")
